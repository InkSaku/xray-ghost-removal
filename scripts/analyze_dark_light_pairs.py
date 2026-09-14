"""Robust causal analysis of the 31 dark/light CR acquisitions.

Stages implemented without relying on InstanceCreationTime:
0. Freeze the evidence boundary; timestamps are audit-only.
1. Test light_(N-1) -> dark_N across analysis choices and causal controls.
2. Test 1--5 frame memory with spatial CV, permutation, and bootstrap.

The default output is one compact JSON. No images or CSVs are generated.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pydicom
from openpyxl import load_workbook
from scipy.optimize import nnls


BLOCK_SIZES = (8, 16, 32, 64)
BASELINE_MODES = ("loo_median", "low_activity", "intercept_only")
MASK_MODES = ("all", "exclude_saturated")
MAX_LAG = 5
N_SPATIAL_NULL = 8
N_BOOTSTRAP = 100
RANDOM_SEED = 20260914


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def block_mean(image: np.ndarray, block: int) -> np.ndarray:
    h = image.shape[0] // block * block
    w = image.shape[1] // block * block
    return image[:h, :w].reshape(h // block, block, w // block, block).mean((1, 3))


def instance_time(ds: pydicom.Dataset) -> str:
    date, time = ds.get("InstanceCreationDate"), ds.get("InstanceCreationTime")
    if not date or not time:
        return ""
    raw = str(time)
    base, _, frac = raw.partition(".")
    try:
        value = f"{date}{base.ljust(6, '0')[:6]}{(frac + '000000')[:6]}"
        return datetime.strptime(value, "%Y%m%d%H%M%S%f").isoformat(sep=" ")
    except ValueError:
        return f"{date} {time}"


def load_exposures(path: Path) -> dict[int, dict[str, float]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    records = {}
    for row in sheet.iter_rows(min_row=4, max_col=4, values_only=True):
        if row[0] is None:
            continue
        index, kv, ma, exposure_ms = int(row[0]), float(row[1]), float(row[2]), float(row[3])
        records[index] = {"kv": kv, "ma": ma, "exposure_time_ms": exposure_ms,
                          "mas": ma * exposure_ms / 1000.0}
    if sorted(records) != list(range(1, 32)):
        raise ValueError("Exposure workbook must contain indices 1 through 31")
    return records


def spatial_folds(shape: tuple[int, int]) -> np.ndarray:
    """Four folds distributed over an 8x8 macro-tile grid."""
    h, w = shape
    rr = np.minimum(np.arange(h) * 8 // h, 7)[:, None]
    cc = np.minimum(np.arange(w) * 8 // w, 7)[None, :]
    return ((rr + 2 * cc) % 4).astype(np.int8)


def trimmed_mask(x: np.ndarray, y: np.ndarray, mask: np.ndarray) -> np.ndarray:
    keep = mask & np.isfinite(x) & np.isfinite(y)
    if keep.sum() < 200:
        return keep
    xv, yv = x[keep], y[keep]
    xlo, xhi = np.percentile(xv, [0.5, 99.5])
    ylo, yhi = np.percentile(yv, [0.5, 99.5])
    return keep & (x >= xlo) & (x <= xhi) & (y >= ylo) & (y <= yhi)


def affine_cv(x: np.ndarray, y: np.ndarray, mask: np.ndarray, folds: np.ndarray) -> dict[str, float]:
    keep = trimmed_mask(x, y, mask)
    if keep.sum() < 200:
        return {"alpha": 0.0, "train_r2": float("nan"), "cv_r2": float("nan"),
                "n_blocks": int(keep.sum())}
    xv, yv = x[keep].astype(float), y[keep].astype(float)
    A = np.column_stack([xv, np.ones_like(xv)])
    (alpha, intercept), *_ = np.linalg.lstsq(A, yv, rcond=None)
    pred = alpha * xv + intercept
    ss = float(np.sum((yv - yv.mean()) ** 2)) + 1e-12
    train_r2 = 1.0 - float(np.sum((yv - pred) ** 2)) / ss
    cv_pred = np.empty_like(yv)
    fv = folds[keep]
    for fold in range(4):
        train, test = fv != fold, fv == fold
        At = np.column_stack([xv[train], np.ones(train.sum())])
        (a, b), *_ = np.linalg.lstsq(At, yv[train], rcond=None)
        cv_pred[test] = a * xv[test] + b
    cv_r2 = 1.0 - float(np.sum((yv - cv_pred) ** 2)) / ss
    return {"alpha": float(alpha), "train_r2": float(train_r2),
            "cv_r2": float(cv_r2), "n_blocks": int(keep.sum())}


def nnls_cv(X: np.ndarray, y: np.ndarray, mask: np.ndarray, folds: np.ndarray) -> dict[str, Any]:
    keep = mask & np.isfinite(y) & np.all(np.isfinite(X), axis=2)
    Xv, yv = X[keep].astype(float), y[keep].astype(float)
    if yv.size < 300:
        return {"alphas": [0.0] * X.shape[2], "train_r2": float("nan"),
                "cv_r2": float("nan"), "n_blocks": int(yv.size)}
    xm, ym = Xv.mean(0), yv.mean()
    coef, _ = nnls(Xv - xm, yv - ym)
    intercept = ym - float(xm @ coef)
    pred = Xv @ coef + intercept
    ss = float(np.sum((yv - ym) ** 2)) + 1e-12
    train_r2 = 1.0 - float(np.sum((yv - pred) ** 2)) / ss
    cv_pred = np.empty_like(yv)
    fv = folds[keep]
    for fold in range(4):
        train, test = fv != fold, fv == fold
        Xt, yt = Xv[train], yv[train]
        xmt, ymt = Xt.mean(0), yt.mean()
        ct, _ = nnls(Xt - xmt, yt - ymt)
        cv_pred[test] = Xv[test] @ ct + ymt - float(xmt @ ct)
    cv_r2 = 1.0 - float(np.sum((yv - cv_pred) ** 2)) / ss
    return {"alphas": [float(v) for v in coef], "train_r2": float(train_r2),
            "cv_r2": float(cv_r2), "n_blocks": int(yv.size)}


def source(light: dict[int, np.ndarray], index: int) -> np.ndarray:
    image = light[index]
    return image - np.percentile(image, 75)


def baseline(dark: dict[int, np.ndarray], target: int, mode: str) -> np.ndarray:
    if mode == "intercept_only":
        return np.zeros_like(dark[target])
    indices = [n for n in range(1, 32) if n != target]
    centered = {n: dark[n] - np.median(dark[n]) for n in indices}
    if mode == "low_activity":
        indices = sorted(indices, key=lambda n: float(np.std(centered[n])))[:10]
    return np.median(np.stack([centered[n] for n in indices]), axis=0)


def spatial_nulls(x: np.ndarray, sat: np.ndarray, seed: int):
    h, w = x.shape
    shifts = [(h // 3, 0), (0, w // 3), (h // 2, w // 2), (h // 4, w // 2)]
    for i, shift in enumerate(shifts):
        yield np.roll(x, shift, (0, 1)), np.roll(sat, shift, (0, 1)), f"roll_{i+1}"
    rng = np.random.default_rng(seed)
    for i in range(N_SPATIAL_NULL - len(shifts)):
        order = rng.permutation(x.size)
        yield x.ravel()[order].reshape(x.shape), sat.ravel()[order].reshape(x.shape), f"shuffle_{i+1}"


def bootstrap_coefficients(X: np.ndarray, y: np.ndarray, mask: np.ndarray, seed: int) -> dict[str, list[float]]:
    h, w, p = X.shape
    rr = np.minimum(np.arange(h) * 8 // h, 7)[:, None]
    cc = np.minimum(np.arange(w) * 8 // w, 7)[None, :]
    groups = (rr * 8 + cc).astype(int)
    valid = [g for g in range(64) if np.any(mask & (groups == g))]
    rng = np.random.default_rng(seed)
    samples = []
    flat_mask, flat_groups = mask.ravel(), groups.ravel()
    for _ in range(N_BOOTSTRAP):
        chosen = rng.choice(valid, len(valid), replace=True)
        idx = np.concatenate([np.flatnonzero(flat_mask & (flat_groups == g)) for g in chosen])
        Xb, yb = X.reshape(-1, p)[idx].astype(float), y.ravel()[idx].astype(float)
        xm, ym = Xb.mean(0), yb.mean()
        coef, _ = nnls(Xb - xm, yb - ym)
        samples.append(coef)
    values = np.stack(samples)
    return {"median": np.median(values, 0).tolist(),
            "ci025": np.percentile(values, 2.5, axis=0).tolist(),
            "ci975": np.percentile(values, 97.5, axis=0).tolist(),
            "positive_frequency": np.mean(values > 1e-10, axis=0).tolist()}


def clean_json(value: Any) -> Any:
    if isinstance(value, dict):
        return {k: clean_json(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean_json(v) for v in value]
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return None if not np.isfinite(value) else float(value)
    if isinstance(value, float) and not np.isfinite(value):
        return None
    return value


def main() -> None:
    parser = argparse.ArgumentParser(description="Robust dark/light afterglow analysis")
    parser.add_argument("--data-dir", default="data/raw/AI修残影例图")
    parser.add_argument("--output", default="outputs/dark_light_analysis_20260913/robustness_results.json")
    args = parser.parse_args()
    data_dir, output = Path(args.data_dir), Path(args.output)
    record_path = data_dir / "拍摄参数记录.xlsx"
    exposures = load_exposures(record_path)

    cache = {b: {"dark": {}, "light": {}, "sat": {}} for b in BLOCK_SIZES}
    audit, acquisition_times = [], 0
    for n in range(1, 32):
        for kind in ("dark", "light"):
            path = data_dir / f"{n}-{kind}.dcm"
            ds = pydicom.dcmread(path)
            image = ds.pixel_array.astype(np.float32)
            acquisition_times += int(bool(ds.get("AcquisitionTime")))
            audit.append({"image": path.name, "rows": int(ds.Rows), "columns": int(ds.Columns),
                          "instance_creation_time": instance_time(ds),
                          "acquisition_time_present": bool(ds.get("AcquisitionTime")),
                          "pixel_mean": float(image.mean()), "pixel_std": float(image.std()),
                          "max_fraction": float(np.mean(image == image.max()))})
            for block in BLOCK_SIZES:
                cache[block][kind][n] = block_mean(image, block)
                if kind == "light":
                    cache[block]["sat"][n] = block_mean((image == image.max()).astype(np.float32), block)

    single_detail, single_pairs = [], []
    for target in range(2, 32):
        prior = target - 1
        pair_rows = []
        for block in BLOCK_SIZES:
            dark, light, sat = cache[block]["dark"], cache[block]["light"], cache[block]["sat"]
            folds, x_true = spatial_folds(dark[target].shape), source(light, prior)
            for base_mode in BASELINE_MODES:
                y = dark[target] - np.median(dark[target]) - baseline(dark, target, base_mode)
                for mask_mode in MASK_MODES:
                    true_mask = np.ones_like(y, bool)
                    if mask_mode == "exclude_saturated":
                        true_mask &= sat[prior] < 0.05
                    fit = affine_cv(x_true, y, true_mask, folds)
                    null, future_scores, h1 = [], [], None
                    for future in range(target, 32):
                        fm = np.ones_like(y, bool)
                        if mask_mode == "exclude_saturated":
                            fm &= sat[future] < 0.05
                        ff = affine_cv(source(light, future), y, fm, folds)
                        if np.isfinite(ff["cv_r2"]):
                            null.append(ff["cv_r2"]); future_scores.append((ff["cv_r2"], future))
                        if future == target:
                            h1 = ff["cv_r2"]
                    for xn, sn, _ in spatial_nulls(x_true, sat[prior], RANDOM_SEED + target * 1000 + block):
                        nm = np.ones_like(y, bool)
                        if mask_mode == "exclude_saturated":
                            nm &= sn < 0.05
                        nf = affine_cv(xn, y, nm, folds)
                        if np.isfinite(nf["cv_r2"]):
                            null.append(nf["cv_r2"])
                    p95, nmed = float(np.percentile(null, 95)), float(np.median(null))
                    ranked = sorted(future_scores + [(fit["cv_r2"], prior)], reverse=True)
                    rank = next(i + 1 for i, (_, idx) in enumerate(ranked) if idx == prior)
                    passed = bool(0 < fit["alpha"] <= 0.02 and fit["cv_r2"] > p95 and rank == 1)
                    row = {"light_index": prior, "dark_index": target, "block_size": block,
                           "baseline_mode": base_mode, "mask_mode": mask_mode, **fit,
                           "h1_future_cv_r2": h1, "null_median_cv_r2": nmed,
                           "null_p95_cv_r2": p95, "delta_cv_r2_vs_null_median": fit["cv_r2"] - nmed,
                           "rank_among_true_and_future": rank, "config_pass": passed}
                    pair_rows.append(row); single_detail.append(row)
        valid = [r for r in pair_rows if np.isfinite(r["cv_r2"])]
        rate = lambda fn: float(np.mean([fn(r) for r in valid]))
        pass_rate = rate(lambda r: r["config_pass"])
        alpha_rate = rate(lambda r: 0 < r["alpha"] <= 0.02)
        cv_rate = rate(lambda r: r["cv_r2"] > 0)
        null_rate = rate(lambda r: r["cv_r2"] > r["null_p95_cv_r2"])
        if pass_rate >= 0.75 and alpha_rate >= 0.90 and cv_rate >= 0.75:
            verdict = "robust_confirmed"
        elif pass_rate > 0 or null_rate >= 0.25:
            verdict = "parameter_sensitive"
        else:
            verdict = "not_detected"
        light_audit = next(r for r in audit if r["image"] == f"{prior}-light.dcm")
        single_pairs.append({"light_index": prior, "dark_index": target, **exposures[prior],
                             "light_max_fraction": light_audit["max_fraction"],
                             "valid_configurations": len(valid), "pass_rate": pass_rate,
                             "alpha_positive_rate": alpha_rate, "cv_positive_rate": cv_rate,
                             "null_exceed_rate": null_rate,
                             "rank1_rate": rate(lambda r: r["rank_among_true_and_future"] == 1),
                             "median_alpha": float(np.median([r["alpha"] for r in valid])),
                             "median_train_r2": float(np.median([r["train_r2"] for r in valid])),
                             "median_cv_r2": float(np.median([r["cv_r2"] for r in valid])),
                             "median_null_cv_r2": float(np.median([r["null_median_cv_r2"] for r in valid])),
                             "verdict": verdict})

    multi_detail, multi_pairs = [], []
    for target in range(2, 32):
        max_lag = min(MAX_LAG, target - 1)
        by_block = {}
        for block in BLOCK_SIZES:
            dark, light, sat = cache[block]["dark"], cache[block]["light"], cache[block]["sat"]
            y = dark[target] - np.median(dark[target]) - baseline(dark, target, "loo_median")
            folds = spatial_folds(y.shape)
            X = np.stack([source(light, target - lag) for lag in range(1, max_lag + 1)], axis=2)
            # Keep one common domain for every lag count. Intersecting the
            # non-saturated regions of up to five differently shaped sources
            # can leave almost no validation blocks and makes incremental R2
            # incomparable. These fits predict from the observed (possibly
            # clipped) DICOM values; alpha bias from clipping is disclosed.
            mask = np.ones_like(y, bool)
            by_block[block] = []
            for n_lags in range(1, max_lag + 1):
                fit = nnls_cv(X[:, :, :n_lags], y, mask, folds)
                row = {"dark_index": target, "block_size": block, "n_lags": n_lags, **fit}
                by_block[block].append(row); multi_detail.append(row)
        medians = {lag: float(np.median([by_block[b][lag - 1]["cv_r2"] for b in BLOCK_SIZES]))
                   for lag in range(1, max_lag + 1)}
        gains = [by_block[b][-1]["cv_r2"] - by_block[b][0]["cv_r2"] for b in BLOCK_SIZES]

        block = 16
        dark, light, sat = cache[block]["dark"], cache[block]["light"], cache[block]["sat"]
        y = dark[target] - np.median(dark[target]) - baseline(dark, target, "loo_median")
        folds = spatial_folds(y.shape)
        X = np.stack([source(light, target - lag) for lag in range(1, max_lag + 1)], axis=2)
        mask = np.ones_like(y, bool)
        fit1, fitall = nnls_cv(X[:, :, :1], y, mask, folds), nnls_cv(X, y, mask, folds)
        real_gain = fitall["cv_r2"] - fit1["cv_r2"]
        perm_gains = []
        if max_lag >= 2:
            rng = np.random.default_rng(RANDOM_SEED + target)
            for _ in range(20):
                shuffled = X.copy()
                for col in range(1, max_lag):
                    order = rng.permutation(X.shape[0] * X.shape[1])
                    shuffled[:, :, col] = shuffled[:, :, col].ravel()[order].reshape(X.shape[:2])
                pf = nnls_cv(shuffled, y, mask, folds)
                perm_gains.append(pf["cv_r2"] - fit1["cv_r2"])
        perm_p95 = float(np.percentile(perm_gains, 95)) if perm_gains else 0.0
        boot = bootstrap_coefficients(X, y, mask, RANDOM_SEED + 1000 + target)
        standardized = X[mask].astype(float)
        standardized = (standardized - standardized.mean(0)) / np.maximum(standardized.std(0), 1e-12)
        condition = float(np.linalg.cond(standardized))
        older_positive = max(boot["positive_frequency"][1:], default=0.0)
        gain_stability = float(np.mean(np.array(gains) > 0))
        supported = bool(max_lag >= 2 and real_gain > perm_p95 and gain_stability >= 0.75 and older_positive >= 0.75)
        multi_pairs.append({"dark_index": target, "max_lag_available": max_lag,
                            "best_lag_by_median_cv_r2": max(medians, key=medians.get),
                            "lag_cv_r2_medians": {str(k): v for k, v in medians.items()},
                            "median_gain_full_vs_lag1_across_blocks": float(np.median(gains)),
                            "gain_positive_block_fraction": gain_stability,
                            "block16_real_gain": real_gain,
                            "block16_permuted_older_lag_gain_p95": perm_p95,
                            "condition_number": condition, "bootstrap": boot,
                            "multi_frame_memory_supported": supported})

    counts = {v: sum(r["verdict"] == v for r in single_pairs)
              for v in ("robust_confirmed", "parameter_sensitive", "not_detected")}
    result = {
        "schema_version": 1, "generated_at": datetime.now().isoformat(timespec="seconds"),
        "stage_0_evidence_boundary": {
            "data_directory": str(data_dir.resolve()), "exposure_workbook": str(record_path.resolve()),
            "exposure_workbook_sha256": sha256(record_path), "image_count": len(audit),
            "acquisition_time_present_count": acquisition_times,
            "recorded_order": "dark_N then light_N; primary causal pair is light_(N-1) -> dark_N",
            "exposure_time_definition": "Workbook column D is light exposure duration in milliseconds.",
            "instance_creation_time_policy": "Audit only; excluded from every fit, control, label, and physical interpretation.",
            "registration": "None; original detector coordinates retained.", "image_audit": audit},
        "configuration": {
            "block_sizes": list(BLOCK_SIZES), "baseline_modes": list(BASELINE_MODES),
            "mask_modes": list(MASK_MODES), "spatial_folds": 4,
            "spatial_nulls_per_configuration": N_SPATIAL_NULL, "bootstrap_repetitions": N_BOOTSTRAP,
            "single_config_pass": "0<alpha<=0.02, CV-R2>null p95, true prior light ranks first against future lights",
            "robust_pair_rule": "pass rate>=0.75, physical alpha rate>=0.90, positive CV-R2 rate>=0.75",
            "multi_memory_rule": "older-lag CV gain>shuffled p95, gain positive in >=75% block sizes, older lag bootstrap-positive >=75%",
            "multi_lag_saturation_policy": "Uses observed DICOM values on a common full domain; coefficients may be biased when a light source is clipped."},
        "stage_1_single_lag": {
            "summary": {"pair_count": len(single_pairs), "configurations_per_pair": 24,
                        "verdict_counts": counts,
                        "median_pass_rate": float(np.median([r["pass_rate"] for r in single_pairs])),
                        "median_alpha": float(np.median([r["median_alpha"] for r in single_pairs])),
                        "median_cv_r2": float(np.median([r["median_cv_r2"] for r in single_pairs])),
                        "median_null_cv_r2": float(np.median([r["median_null_cv_r2"] for r in single_pairs]))},
            "pair_summary": single_pairs, "configuration_detail": single_detail},
        "stage_2_multi_lag": {
            "summary": {"dark_count": len(multi_pairs),
                        "multi_frame_memory_supported_count": sum(r["multi_frame_memory_supported"] for r in multi_pairs),
                        "median_full_vs_lag1_cv_gain": float(np.median([r["median_gain_full_vs_lag1_across_blocks"] for r in multi_pairs]))},
            "dark_summary": multi_pairs, "fit_detail": multi_detail},
        "interpretation_limits": ["No time-decay constant or half-life is estimated.",
                                  "InstanceCreationTime is audit-only.",
                                  "Exposure association is outside stages 0-2.",
                                  "Dark images are calibration targets, not clean references for exposed light images."]}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(clean_json(result), ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps(clean_json({"stage_1": result["stage_1_single_lag"]["summary"],
                                 "stage_2": result["stage_2_multi_lag"]["summary"],
                                 "output": str(output.resolve())}), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
