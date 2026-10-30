"""Exposure workbook readers for the existing 31-group dataset."""
from pathlib import Path
from openpyxl import load_workbook
from .provenance import sha256

def load_exposures(path: Path) -> dict[int, dict[str, float]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    records = {}
    for row in sheet.iter_rows(min_row=4, max_col=4, values_only=True):
        if row[0] is None:
            continue
        index, kv, ma, exposure_ms = int(row[0]), float(row[1]), float(row[2]), float(row[3])
        records[index] = {"kv": kv, "ma": ma, "exposure_time_ms": exposure_ms,
                          "mas": ma * exposure_ms / 1000.0}
    if sorted(records) != list(range(1, 32)):
        raise ValueError("Exposure workbook must contain indices 1 through 31")
    return records


def workbook_record(path: Path) -> dict:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = [dict(index=int(a), kv=float(b), ma=float(c), ms=float(d),
                 mas=float(c)*float(d)/1000, row=i)
            for i, (a, b, c, d) in enumerate(sheet.iter_rows(min_row=4, max_col=4, values_only=True), 4)
            if a is not None]
    return dict(path=str(path), sha256=sha256(path), note=sheet["A1"].value, rows=rows)
